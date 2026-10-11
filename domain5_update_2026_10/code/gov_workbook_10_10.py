"""Government workbook (10.10.2026): zero-dose estimates 2026 by LGA, Method 1 vs Method 2 (Bayesian small-area estimation),
with contextual profiles, each LGA's dominant barrier and candidate intervention components, priority certainty and the
external-validation comparison. Labels describe only what the input indicators measure."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.views import Selection

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
OUT = ROOT / "government_workbook"; OUT.mkdir(exist_ok=True)
NAVY, STEEL, LIGHT, M1F, M2F = "1F3B57", "2E6E8E", "EEF2F6", "DCE9F5", "FBE5D6"
thin = Side(style="thin", color="D5DCE3")
M1 = "Method 1: Bayesian hierarchical model with DHIS2-calibrated LGA allocation"
M2 = "Method 2: Bayesian small-area estimation (SAE)"

lg = pd.read_csv(RES / "lga_both_methods_774.csv")
st = pd.read_csv(RES / "state_both_methods.csv")
zn = pd.read_csv(RES / "national_zone_both_methods.csv")
val = pd.read_csv(RES / "validation_both_methods.csv")
par = pd.read_csv(RES / "pareto_scenarios.csv")
tk = pd.read_csv(RES / "topk_shares.csv")
ar = pd.read_csv(RES / "archetype_both_methods.csv")
R = json.loads((RES / "results.json").read_text())
EV = ROOT.parents[1] / "nphcda-zerodose-git" / "data" / "sample" / "two_methods" / "evidence"
PLAB = {1: "Deprived northern rural", 2: "Low maternal care, high conflict exposure", 3: "Geographically isolated",
        4: "Near-average", 5: "Relatively advantaged"}
lg = lg.loc[:, ~lg.columns.duplicated()]
lg["archetype_type"] = lg["archetype"].map(PLAB)
ar["archetype_type"] = ar["archetype"].map(PLAB)
bar = pd.read_csv(EV / "lga_barriers_774.csv")
lg = lg.merge(bar[["lga_uid", "dominant_barrier", "barriers_flagged", "n_barriers", "candidate_components", "membership"]],
              on="lga_uid", how="left")
for c in ("dominant_barrier", "barriers_flagged"):
    lg[c] = lg[c].str.replace("Programme (dropout)", "DTP1-3 dropout", regex=False)
lg["m2_certainty"] = np.select([lg.m2_p_top155 >= 0.9, lg.m2_p_top155 <= 0.1], ["Near-certain", "Unlikely"], "Uncertain")
COMP = pd.read_csv(EV / "validation_comparators.csv")
RULES = pd.read_csv(EV / "assignment_rules.csv")
BMAT = pd.read_csv(EV / "barrier_matrix.csv")

wb = Workbook()


def header(ws, row, cols, fill=NAVY):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10); cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); cell.border = Border(bottom=thin)


def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=15, color=NAVY)
    if sub:
        ws["A2"] = sub; ws["A2"].font = Font(italic=True, size=9.5, color="5B6B79")


def write_table(ws, df, start_row, fmts=None, group_fills=None, name=None, widths=None, height=30):
    header(ws, start_row, list(df.columns))
    ws.row_dimensions[start_row].height = height
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10); c.border = Border(bottom=thin)
            col = df.columns[j - 1]
            if fmts and col in fmts:
                c.number_format = fmts[col]
            if group_fills:
                for prefix, f in group_fills.items():
                    if col.startswith(prefix):
                        c.fill = PatternFill("solid", fgColor=f)
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(col, max(11, min(42, len(str(col)) * 0.9)))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=4 if len(df.columns) > 6 else 2)
    if name:
        ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{start_row}:{start_row}"
    return start_row + len(df)


# ---------------------------------------------------------------- 1. About
ws = wb.active; ws.title = "About"
title(ws, "Zero-dose children in Nigeria, 2026: estimates for all local government areas",
      "National Primary Health Care Development Agency (NPHCDA). Model estimates prepared by the CIDRE and Quantium Insights LLC consortium, with UNICEF and Gavi.")
rows = [
    ("What this workbook contains", "Model estimates of zero-dose children (children aged 12-23 months who have not received the first dose of the pentavalent vaccine, Penta1/DTP1) for 2026, for every local government area (LGA), state and zone, from two Bayesian methods, shown side by side."),
    ("Method 1", M1 + ". A Bayesian hierarchical Beta regression estimates each state's zero-dose rate from the Nigeria Demographic and Health Surveys (NDHS) 2008, 2013, 2018 and 2023-24, with a DHIS2 Penta1 trend term, and forecasts it to 2026. Each state's rate is then distributed to its LGAs using each LGA's share of Penta1 doses in DHIS2 and its population. Estimates for 773 LGAs; Guzamala (Borno) is not estimated because it reported no Penta1 doses in 2021-2024, the years used for the allocation (partial reporting began in 2025)."),
    ("Method 2", M2 + ". Each LGA has its own zero-dose rate. The population-weighted average of a state's LGA rates must agree with that state's NDHS results (allowing for survey sampling error). Differences between LGAs are informed by local conditions (maternal care, water, wealth, travel time to a facility, conflict and poverty) and by neighbouring LGAs (spatial smoothing). Estimates for all 774 LGAs, with full uncertainty."),
    ("Population base", "Children aged 12-23 months = state under-five population (2024 projection) divided by five, shared among LGAs by National Population Commission (NPC) 2022 LGA population."),
    ("Uncertainty", "Both methods give a 95% interval. For Method 1 it reflects uncertainty in the state model; for Method 2 it is a full Bayesian credible interval for each LGA. Neither interval includes error in the population figures."),
    ("Priority probability", "The share of model simulations in which the LGA is among the 155 highest-burden LGAs (top 20% of LGAs). Values near 1 mean the LGA is consistently a priority; values between 0.1 and 0.9 mean its position is uncertain."),
    ("Agreement with independent data", "Both methods were checked against the Nigeria mini Demographic and Health Survey (NmDHS) 2025-26 at state level (not used to build either model) and the IHME Local Burden of Disease DTP1 map (2018) at LGA level. See the 'Validation' sheet."),
    ("How to read the two methods", "Both methods agree closely on national, zone and state totals. They differ mainly in how zero-dose children are distributed among the LGAs of a state: Method 1 follows routine Penta1 reporting patterns; Method 2 follows survey levels and local structural conditions. LGAs that rank high under both methods are the most robust priorities (see 'Priority lists')."),
    ("Contextual profiles", "Five groups of LGAs from Ward clustering of 15 contextual indicators in six domains (socioeconomic, maternal care, nutrition, access, insecurity, DTP1-3 dropout): deprived northern rural; low maternal care, high conflict exposure; geographically isolated; near-average; relatively advantaged. Labels describe only what the indicators measure; mobility, migration, riverine settlement and informal settlements are not measured."),
    ("Barriers and candidate components", "An LGA barrier is flagged when its domain score is in the national top quartile (travel time: top decile). Candidate components follow the flagged barriers (see 'Profiles and packages') and must be validated locally through microplanning. They are not measured intervention effects."),
    ("Sheets", "National and zone | States | LGA comparison (all 774) | Priority lists | Burden concentration | Validation | Profiles and packages | Definitions"),
    ("Data sources", "NDHS 2008-2024 (DHS Program); NPHCDA DHIS2 routine immunization data 2021-2025; NPC 2022 population projections; GRID3 administrative boundaries; DHS Spatial Data Repository (antenatal care, facility delivery, improved water, 2018); Meta Relative Wealth Index; Weiss et al. travel time to healthcare; ACLED conflict events 2021-2024; poverty surface; NmDHS 2025-26 and IHME Local Burden of Disease DTP1 (validation only)."),
    ("Status", "Model estimates for programme planning. Figures are estimates, not counts, and should be read with their intervals."),
]
for i, (k, v) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True, color=NAVY, size=10.5)
    c = ws.cell(row=i, column=2, value=v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(size=10.5)
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    ws.row_dimensions[i].height = max(30, 15 * (len(v) // 120 + 1))
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 130
r0 = 17
ws.cell(row=r0, column=1, value="Headline figures").font = Font(bold=True, size=12, color=NAVY)
heads = [("", "Method 1", "Method 2 (SAE)"),
         ("LGAs estimated", 773, 774),
         ("Zero-dose children, 2026", R["m1"]["national"], R["m2"]["national"]),
         ("95% interval, lower", R["m1"]["national_lo95"], R["m2"]["national_lo95"]),
         ("95% interval, upper", R["m1"]["national_hi95"], R["m2"]["national_hi95"]),
         ("LGAs holding 50% of zero-dose children", R["m1"]["pareto"]["50"], R["m2"]["pareto"]["50"]),
         ("LGAs holding 60% of zero-dose children", R["m1"]["pareto"]["60"], R["m2"]["pareto"]["60"]),
         ("LGAs holding 80% of zero-dose children", R["m1"]["pareto"]["80"], R["m2"]["pareto"]["80"]),
         ("Share of zero-dose children in top 155 LGAs (20%)", R["m1"]["top155_share"] / 100, R["m2"]["top155_share"] / 100)]
for i, rr in enumerate(heads, r0 + 1):
    for j, v in enumerate(rr, 1):
        c = ws.cell(row=i, column=j + 0 if j == 1 else j, value=v)
    ws.cell(row=i, column=1).font = Font(bold=(i == r0 + 1), color=NAVY)
for i in range(r0 + 1, r0 + len(heads) + 1):
    for j in (2, 3):
        c = ws.cell(row=i, column=j); c.alignment = Alignment(horizontal="right")
        c.number_format = "0.0%" if "Share" in str(ws.cell(row=i, column=1).value) else "#,##0"
        if i == r0 + 1:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
ws.column_dimensions["C"].width = 18
ws.cell(row=r0 + 1, column=1).fill = PatternFill("solid", fgColor=NAVY); ws.cell(row=r0 + 1, column=1).font = Font(bold=True, color="FFFFFF")
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- 2. National and zone
ws = wb.create_sheet("National and zone")
title(ws, "Zero-dose children by zone, 2026: Method 1 and Method 2", "Model estimates with 95% intervals. Shares are of each method's national total.")
z = zn.rename(columns={"area": "Area", "lgas_total": "LGAs", "m1_lgas": "LGAs estimated (M1)", "cohort_12_23m": "Children 12-23 months",
                       "m1_children": "M1 zero-dose children", "m1_lo95": "M1 lower 95%", "m1_hi95": "M1 upper 95%",
                       "m1_rate": "M1 zero-dose rate (%)", "m1_share": "M1 share of national (%)",
                       "m2_children": "M2 zero-dose children", "m2_lo95": "M2 lower 95%", "m2_hi95": "M2 upper 95%",
                       "m2_rate": "M2 zero-dose rate (%)", "m2_share": "M2 share of national (%)",
                       "difference": "Difference (M2 - M1)", "difference_pct": "Difference (%)"})
z = z[["Area", "LGAs", "LGAs estimated (M1)", "Children 12-23 months", "M1 zero-dose children", "M1 lower 95%", "M1 upper 95%",
       "M1 zero-dose rate (%)", "M1 share of national (%)", "M2 zero-dose children", "M2 lower 95%", "M2 upper 95%",
       "M2 zero-dose rate (%)", "M2 share of national (%)", "Difference (M2 - M1)", "Difference (%)"]]
f = {c: "#,##0" for c in z.columns if "children" in c.lower() or "95%" in c or "Difference (M2" in c}
f.update({c: "0.0" for c in z.columns if "(%)" in c})
end = write_table(ws, z, 4, f, {"M1": M1F, "M2": M2F}, "ZoneTable")
for col in range(1, len(z.columns) + 1):
    ws.cell(row=5, column=col).font = Font(bold=True, size=10)

# ---------------------------------------------------------------- 3. States
ws = wb.create_sheet("States")
title(ws, "Zero-dose estimates by state, 2026: Method 1 and Method 2",
      "Rates are state zero-dose percentages; children are model estimates with 95% intervals. Survey columns are measured values for reference.")
s = st.rename(columns={"state": "State", "zone": "Zone", "lgas": "LGAs", "m1_lgas": "LGAs estimated (M1)",
                       "ndhs_2024": "NDHS 2023-24 zero-dose (%)", "nmdhs_2025_26": "NmDHS 2025-26 zero-dose (%)",
                       "m1_rate": "M1 rate 2026 (%)", "m1_rate_lo95": "M1 rate lower 95%", "m1_rate_hi95": "M1 rate upper 95%",
                       "m1_children": "M1 zero-dose children", "m1_children_lo95": "M1 children lower 95%", "m1_children_hi95": "M1 children upper 95%",
                       "m2_rate": "M2 rate 2026 (%)", "m2_rate_lo95": "M2 rate lower 95%", "m2_rate_hi95": "M2 rate upper 95%",
                       "m2_children": "M2 zero-dose children", "m2_children_lo95": "M2 children lower 95%", "m2_children_hi95": "M2 children upper 95%",
                       "cohort_12_23m": "Children 12-23 months", "m1_lga_rate_range": "M1 LGA rate range (%)",
                       "m2_lga_rate_range": "M2 LGA rate range (%)", "difference": "Difference (M2 - M1)", "difference_pct": "Difference (%)"})
s = s[["State", "Zone", "LGAs", "LGAs estimated (M1)", "Children 12-23 months", "NDHS 2023-24 zero-dose (%)", "NmDHS 2025-26 zero-dose (%)",
       "M1 rate 2026 (%)", "M1 rate lower 95%", "M1 rate upper 95%", "M1 zero-dose children", "M1 children lower 95%", "M1 children upper 95%", "M1 LGA rate range (%)",
       "M2 rate 2026 (%)", "M2 rate lower 95%", "M2 rate upper 95%", "M2 zero-dose children", "M2 children lower 95%", "M2 children upper 95%", "M2 LGA rate range (%)",
       "Difference (M2 - M1)", "Difference (%)"]]
f = {c: "#,##0" for c in s.columns if "children" in c.lower() or "Difference (M2" in c}
f.update({c: "0.0" for c in s.columns if "(%)" in c or "rate lower" in c or "rate upper" in c})
end = write_table(ws, s, 4, f, {"M1": M1F, "M2": M2F}, "StateTable")
for col in ("K", "R"):
    ws.conditional_formatting.add(f"{col}5:{col}{end}", DataBarRule(start_type="min", end_type="max", color="5B9BD5" if col == "K" else "ED7D31"))

# ---------------------------------------------------------------- 4. LGA comparison
ws = wb.create_sheet("LGA comparison (all 774)")
title(ws, "Zero-dose estimates for every LGA, 2026: Method 1 and Method 2",
      "Sorted by Method 2 burden rank. Use the filter arrows to select a state or zone. Priority probability = chance the LGA is among the 155 highest-burden LGAs.")
L = lg.copy()
L["m1_cum"] = np.nan
o1 = L.dropna(subset=["m1_children"]).sort_values("m1_rank")
L.loc[o1.index, "m1_cum"] = o1.m1_children.cumsum() / o1.m1_children.sum() * 100
o2 = L.sort_values("m2_rank"); L.loc[o2.index, "m2_cum"] = o2.m2_children.cumsum() / o2.m2_children.sum() * 100
band = lambda c: np.where(c.isna(), "", np.where(c <= 50, "A: first 50%", np.where(c <= 80, "B: 50-80%", "C: last 20%")))  # noqa: E731
L["m1_band"] = band(L.m1_cum); L["m2_band"] = band(L.m2_cum)
L["agree"] = np.select([(L.m1_rank <= 155) & (L.m2_rank <= 155), (L.m2_rank <= 155), (L.m1_rank <= 155)],
                       ["Top 155 under both methods", "Top 155 under Method 2 only", "Top 155 under Method 1 only"], "")
L["m1_note"] = np.where(L.m1_children.isna(), "Not estimated: no Penta1 doses reported 2021-2024 (partial reporting from 2025)",
                        np.where(L.m1_capped_99 == True, "Method 1 rate at its 99% ceiling", ""))  # noqa: E712
cols = {"zone": "Zone", "state": "State", "lga_clean": "LGA", "cohort_12_23m": "Children 12-23 months",
        "m1_rate": "M1 rate (%)", "m1_rate_lo95": "M1 rate lower 95%", "m1_rate_hi95": "M1 rate upper 95%",
        "m1_children": "M1 zero-dose children", "m1_children_lo95": "M1 children lower 95%", "m1_children_hi95": "M1 children upper 95%",
        "m1_rank": "M1 national rank (of 773)", "m1_p_top155": "M1 priority probability", "m1_band": "M1 burden band",
        "m2_rate": "M2 rate (%)", "m2_rate_lo95": "M2 rate lower 95%", "m2_rate_hi95": "M2 rate upper 95%",
        "m2_children": "M2 zero-dose children", "m2_children_lo95": "M2 children lower 95%", "m2_children_hi95": "M2 children upper 95%",
        "m2_rank": "M2 national rank (of 774)", "m2_rank_lo95": "M2 rank lower 95%", "m2_rank_hi95": "M2 rank upper 95%",
        "m2_p_top155": "M2 priority probability", "m2_band": "M2 burden band", "agree": "Top-155 agreement",
        "archetype": "Profile", "archetype_type": "Contextual profile", "membership": "Profile membership",
        "dominant_barrier": "Dominant barrier", "barriers_flagged": "Flagged barriers",
        "candidate_components": "Candidate components (local validation needed)", "m1_note": "Note"}
Lw = L.sort_values("m2_rank")[list(cols)].rename(columns=cols)
Lw["Difference in children (M2 - M1)"] = Lw["M2 zero-dose children"] - Lw["M1 zero-dose children"]
f = {c: "#,##0" for c in Lw.columns if "children" in c.lower() or "rank" in c.lower() or "Difference" in c}
f.update({c: "0.0" for c in Lw.columns if "rate" in c.lower()})
f.update({c: "0.00" for c in Lw.columns if "probability" in c})
end = write_table(ws, Lw, 4, f, {"M1": M1F, "M2": M2F}, "LGATable", widths={"LGA": 24, "Note": 40, "Contextual profile": 30, "Top-155 agreement": 26, "Flagged barriers": 30,
                                                                           "Candidate components (local validation needed)": 48})
cl = {c: get_column_letter(i + 1) for i, c in enumerate(Lw.columns)}
for c, color in [("M1 priority probability", "5B9BD5"), ("M2 priority probability", "ED7D31")]:
    ws.conditional_formatting.add(f"{cl[c]}5:{cl[c]}{end}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                          end_type="num", end_value=1, end_color=color))
for c, color in [("M1 zero-dose children", "5B9BD5"), ("M2 zero-dose children", "ED7D31")]:
    ws.conditional_formatting.add(f"{cl[c]}5:{cl[c]}{end}", DataBarRule(start_type="num", start_value=0, end_type="max", color=color))

# ---------------------------------------------------------------- 5. Priority lists
ws = wb.create_sheet("Priority lists")
title(ws, "Priority LGAs: the 155 highest-burden LGAs (top 20%) under each method",
      "LGAs in the top 155 under both methods are the most robust priorities. Probability = chance of being in the top 155 under each method.")
P = L[(L.m1_rank <= 155) | (L.m2_rank <= 155)].copy()
P["order"] = np.where(P.agree == "Top 155 under both methods", 0, np.where(P.agree.str.contains("Method 2"), 1, 2))
P = P.sort_values(["order", "m2_rank"])
pc = {"agree": "Top-155 agreement", "zone": "Zone", "state": "State", "lga_clean": "LGA", "m1_rank": "M1 rank", "m1_children": "M1 zero-dose children",
      "m1_p_top155": "M1 priority probability", "m2_rank": "M2 rank", "m2_children": "M2 zero-dose children", "m2_p_top155": "M2 priority probability",
      "m2_certainty": "M2 priority certainty", "archetype_type": "Contextual profile", "dominant_barrier": "Dominant barrier",
      "candidate_components": "Candidate components (local validation needed)"}
Pw = P[list(pc)].rename(columns=pc)
f = {"M1 rank": "#,##0", "M2 rank": "#,##0", "M1 zero-dose children": "#,##0", "M2 zero-dose children": "#,##0",
     "M1 priority probability": "0.00", "M2 priority probability": "0.00"}
end = write_table(ws, Pw, 4, f, {"M1": M1F, "M2": M2F}, "PriorityTable", widths={"Top-155 agreement": 28, "LGA": 24, "Contextual profile": 30,
                                                                              "Candidate components (local validation needed)": 48})
cnt = Pw["Top-155 agreement"].value_counts()
ws.cell(row=3, column=1, value=f"Top 155 under both methods: {cnt.get('Top 155 under both methods', 0)} LGAs; Method 2 only: {cnt.get('Top 155 under Method 2 only', 0)}; Method 1 only: {cnt.get('Top 155 under Method 1 only', 0)}.").font = Font(bold=True, color=STEEL)

# ---------------------------------------------------------------- 6. Burden concentration
ws = wb.create_sheet("Burden concentration")
title(ws, "How many LGAs hold 50%, 60% and 80% of zero-dose children", "LGAs ranked by each method's estimated burden. Interval = share captured by that list across model simulations.")
pp = par.rename(columns={"share_of_burden": "Share of zero-dose children (%)", "m1_lgas": "M1 LGAs needed", "m1_pct_of_lgas": "M1 % of LGAs",
                         "m1_list_share_lo95": "M1 list share lower 95%", "m1_list_share_hi95": "M1 list share upper 95%",
                         "m2_lgas": "M2 LGAs needed", "m2_pct_of_lgas": "M2 % of LGAs",
                         "m2_list_share_lo95": "M2 list share lower 95%", "m2_list_share_hi95": "M2 list share upper 95%"})
f = {c: "0.0" for c in pp.columns if "%" in c or "share" in c}; f.update({"M1 LGAs needed": "0", "M2 LGAs needed": "0"})
end = write_table(ws, pp, 4, f, {"M1": M1F, "M2": M2F}, "ParetoTable")
tt = tk.rename(columns={"top_k": "Top LGAs", "m1_share": "M1 share (%)", "m1_lo95": "M1 lower 95%", "m1_hi95": "M1 upper 95%",
                        "m2_share": "M2 share (%)", "m2_lo95": "M2 lower 95%", "m2_hi95": "M2 upper 95%",
                        "m1_near_certain": "M1 LGAs with priority probability >= 0.9", "m2_near_certain": "M2 LGAs with priority probability >= 0.9"})
ws.cell(row=end + 3, column=1, value="Share of zero-dose children in the top-ranked LGAs").font = Font(bold=True, size=12, color=NAVY)
write_table(ws, tt, end + 4, {c: "0.0" for c in tt.columns if "%" in c}, {"M1": M1F, "M2": M2F}, "TopKTable")

# ---------------------------------------------------------------- 7. Validation
ws = wb.create_sheet("Validation")
title(ws, "Agreement with independent data (Spearman rank correlation)",
      "NmDHS 2025-26 was not used to build either model. IHME 2018 is eight years earlier than the 2026 estimates and tests the geographic pattern, not 2026 levels.")
V = val[val.method.isin(["Method 1", "Method 2 (SAE)"])].copy()
V["check"] = V.check.replace({"NmDHS 2025-26, 37 states (state 2026 estimate vs survey)": "States vs NmDHS 2025-26 (37 states)",
                              "IHME DTP1 2018, LGA overall": "LGAs vs IHME 2018, national ranking",
                              "IHME DTP1 2018, LGA within state (pooled within-state ranks)": "LGAs vs IHME 2018, ranking within each state"})
V = V.rename(columns={"method": "Method", "check": "Comparison", "rho": "Spearman rho", "ci_lo": "95% CI lower", "ci_hi": "95% CI upper",
                      "p_value": "p-value", "n": "n", "MAE_pp": "Mean absolute error (percentage points)", "bias_pp": "Mean bias (percentage points)",
                      "states_positive": "States with positive within-state agreement", "states_tested": "States tested"})
V = V[["Comparison", "Method", "n", "Spearman rho", "95% CI lower", "95% CI upper", "p-value", "Mean absolute error (percentage points)",
       "Mean bias (percentage points)", "States with positive within-state agreement", "States tested"]].sort_values(["Comparison", "Method"])
V["Method"] = V.Method.replace({"Method 2 (SAE)": "Method 2 (SAE)"})
end = write_table(ws, V, 4, {"Spearman rho": "0.00", "95% CI lower": "0.00", "95% CI upper": "0.00", "p-value": "0.0000",
                             "Mean absolute error (percentage points)": "0.0", "Mean bias (percentage points)": "0.0"},
                  name="ValidationTable", widths={"Comparison": 44})
ws.cell(row=end + 2, column=1, value="p-values below 0.0001 are shown as 0.0000. Confidence intervals from 2,000 bootstrap resamples.").font = Font(italic=True, size=9, color="5B6B79")

ws.cell(row=end + 4, column=1, value="Accuracy against the NmDHS 2025-26 and survey-round holdouts: Method 2 and alternative approaches").font = Font(bold=True, size=12, color=NAVY)
C2 = COMP.copy()
C2["check"] = C2["check"].replace({"NmDHS 2025-26": "States vs NmDHS 2025-26 (37 states)", "Predict 2018": "Holdout: predict 2018 from 2008-2013",
                                   "Predict 2024": "Holdout: predict 2024 from 2008-2018"})
C2.loc[C2.approach.isin(["Gradient boosting", "Routine data only"]), "CRPS"] = np.nan
C2 = C2.rename(columns={"check": "Check", "approach": "Approach", "MAE": "Mean absolute error (points)", "bias": "Mean bias (points)",
                        "spearman_rho": "Spearman rho", "dMAE_vs_SAE": "Extra error over Method 2 (points)",
                        "dMAE_lo": "Extra error lower 95%", "dMAE_hi": "Extra error upper 95%"})
C2 = C2[["Check", "Approach", "Mean absolute error (points)", "Mean bias (points)", "Spearman rho", "CRPS",
         "Extra error over Method 2 (points)", "Extra error lower 95%", "Extra error upper 95%"]].sort_values(["Check", "Mean absolute error (points)"])
fm = {c: "0.0" for c in C2.columns if "points" in c or c == "CRPS" or "95%" in c}
fm["Spearman rho"] = "0.00"
end2 = write_table(ws, C2, end + 5, fm, name="AccuracyTable", widths={"Check": 44, "Approach": 36})
ws.cell(row=end2 + 2, column=1, value="Calibration (design effect 2): the 95% predictive interval contained the NmDHS value in 37 of 37 states and the 50% interval in 24 of 37. Method 2 priority certainty: 40 LGAs near-certain to be in the top 155 (probability 0.9 or more), 217 uncertain, 517 unlikely. CRPS = continuous ranked probability score, blank for approaches without a predictive distribution.").font = Font(italic=True, size=9, color="5B6B79")
ws.freeze_panes = "A5"
ws.sheet_view.selection = [Selection(pane="bottomLeft", activeCell="A5", sqref="A5")]

# ---------------------------------------------------------------- 8. Archetypes
ws = wb.create_sheet("Profiles and packages")
title(ws, "Zero-dose children by contextual profile, 2026, and the barrier-to-intervention mapping",
      "Profiles from Ward clustering of 15 contextual indicators for all 774 LGAs; labels describe only what the indicators measure.")
A = ar.rename(columns={"archetype": "Profile", "archetype_type": "Name", "lgas": "LGAs", "m1_lgas": "LGAs estimated (M1)",
                       "m1_children": "M1 zero-dose children", "m1_share": "M1 share (%)", "m1_mean_rate": "M1 mean LGA rate (%)",
                       "m2_children": "M2 zero-dose children", "m2_share": "M2 share (%)", "m2_mean_rate": "M2 mean LGA rate (%)",
                       "cohort": "Children 12-23 months", "cohort_share": "Share of children 12-23 months (%)"})
A = A[["Profile", "Name", "LGAs", "LGAs estimated (M1)", "Children 12-23 months", "Share of children 12-23 months (%)", "M1 zero-dose children",
       "M1 share (%)", "M1 mean LGA rate (%)", "M2 zero-dose children", "M2 share (%)", "M2 mean LGA rate (%)"]]
write_table(ws, A, 4, {c: "#,##0" for c in A.columns if "children" in c.lower()} | {c: "0.0" for c in A.columns if "(%)" in c},
            {"M1": M1F, "M2": M2F}, "ProfileTable", widths={"Name": 36})

r = 4 + len(A) + 3
ws.cell(row=r - 1, column=1, value="Barrier-to-intervention mapping (applied to each LGA)").font = Font(bold=True, size=12, color=NAVY)
B = BMAT.rename(columns={"barrier": "Barrier", "indicator": "Indicators", "rule": "Flagging rule", "components": "Candidate components",
                         "evidence": "Evidence or guideline", "local_validation": "Local validation required"})
B["Barrier"] = B["Barrier"].str.replace("Programme (dropout)", "DTP1-3 dropout", regex=False)
end = write_table(ws, B, r, name="BarrierTable")
for j, w in enumerate([26, 36, 34, 40, 40, 36], 1):
    ws.column_dimensions[get_column_letter(j)].width = max(w, ws.column_dimensions[get_column_letter(j)].width or 0)
    for i in range(r + 1, end + 1):
        ws.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
r = end + 3
ws.cell(row=r - 1, column=1, value="Agreement of package-assignment rules with each LGA barrier package (Jaccard; 1 = identical)").font = Font(bold=True, size=12, color=NAVY)
RR = RULES.rename(columns={"assignment_rule": "Assignment rule", "mean_jaccard_vs_lga": "Mean agreement (LGAs)",
                           "burden_weighted_jaccard": "Agreement weighted by modelled zero-dose children"}).sort_values(
    "Agreement weighted by modelled zero-dose children", ascending=False)
end = write_table(ws, RR, r, {"Mean agreement (LGAs)": "0.00", "Agreement weighted by modelled zero-dose children": "0.00"}, name="RuleTable")
ws.cell(row=end + 2, column=1, value="State packages refined with each LGA's flagged barriers agree best; one national package agrees least. A common core of fixed-site and outreach routine immunization with microplanning applies to every LGA. Agreement is with documented rules, not measured intervention effects.").font = Font(italic=True, size=9, color="5B6B79")
ws.freeze_panes = "A5"
ws.sheet_view.selection = [Selection(pane="bottomLeft", activeCell="A5", sqref="A5")]

# ---------------------------------------------------------------- 9. Definitions
ws = wb.create_sheet("Definitions")
title(ws, "Column definitions")
defs = [("M1 / M2", "Method 1 / Method 2 (see the About sheet)."),
        ("Zero-dose rate (%)", "Estimated percentage of children aged 12-23 months who have not received Penta1 (DTP1), 2026."),
        ("Zero-dose children", "Estimated number of zero-dose children aged 12-23 months, 2026 = rate x children 12-23 months."),
        ("Lower / upper 95%", "95% interval. Method 1: state-model uncertainty carried to the LGA. Method 2: Bayesian 95% credible interval."),
        ("National rank", "Rank among LGAs by estimated zero-dose children (1 = largest)."),
        ("M2 rank lower / upper 95%", "Range of ranks the LGA takes in 95% of Method 2 simulations."),
        ("Priority probability", "Share of model simulations in which the LGA is among the 155 highest-burden LGAs (top 20%)."),
        ("Burden band", "A = LGAs that together hold the first 50% of zero-dose children; B = the next 30% (to 80%); C = the remaining 20%."),
        ("Top-155 agreement", "Whether the LGA is in the top 155 under both methods, one method, or neither."),
        ("Method 1 rate at its 99% ceiling", "In Method 1 the allocation caps an LGA's rate at 99%; these LGAs should be checked against local data."),
        ("Contextual profile", "Group of LGAs from Ward clustering of 15 contextual indicators (1 deprived northern rural; 2 low maternal care, high conflict exposure; 3 geographically isolated; 4 near-average; 5 relatively advantaged). Labels describe only what the indicators measure."),
        ("Profile membership", "Clear: bootstrap assignment frequency 0.8 or more and centroid distance ratio below 0.8; mixed: frequency below 0.5 or ratio 0.9 or more; intermediate: all others. Stability and affinity measures, not probabilities."),
        ("Dominant barrier", "The contextual domain with the highest score for the LGA, if at least +0.5 standard deviations above the national mean."),
        ("Flagged barriers", "Domains whose score is in the national top quartile (travel time: top decile)."),
        ("Candidate components", "Intervention components matched to the flagged barriers (see Profiles and packages). To be validated locally through microplanning before adoption."),
        ("M2 priority certainty", "Near-certain: Method 2 priority probability 0.9 or more; unlikely: 0.1 or less; uncertain: in between. Verify uncertain LGAs before full investment."),
        ("Mean absolute error / CRPS", "Average absolute difference from the survey value in percentage points; continuous ranked probability score, which rewards both accuracy and sharp intervals (lower is better).")]
header(ws, 3, ["Term", "Definition"])
for i, (k, v) in enumerate(defs, 4):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True, color=NAVY)
    c = ws.cell(row=i, column=2, value=v); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 120

out = OUT / "NPHCDA_ZeroDose_LGA_Estimates_2026_Method1_vs_SAE_10.10.2026.xlsx"
wb.save(out)
print(out)
