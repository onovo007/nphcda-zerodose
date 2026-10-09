"""Domain 1 government workbook: LGA early-warning forecasts for 773 LGAs, worklist, summaries, Guzamala verification."""
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "corrected"
DATA = ROOT.parent / "All I Need" / "03_Datasets" / "datasets"
WB = ROOT / "government_workbook" / "NPHCDA_Domain1_Antigen_EarlyWarning_LGA_2026.xlsx"
WB.parent.mkdir(exist_ok=True)
ANT = {"BCG": "bcg_count", "Penta1": "penta_1_count", "Penta3": "penta_3_count", "Measles1": "measles_1_count"}
NAVY = "1F3B57"


def clean(s):
    s = str(s).strip()
    for suf in (" Local Government Area", " LGA"):
        if s.lower().endswith(suf.lower()):
            s = s[: -len(suf)]
    p = s.split(None, 1)
    return (p[1] if len(p) > 1 and len(p[0]) == 2 and p[0].islower() else s).strip()


d = pd.read_csv(DATA / "dhis2_data_all_states.csv", thousands=",", low_memory=False)
d["ds"] = pd.to_datetime(d.period.astype(str).str.strip(), format="%b-%y")
l = d.groupby(["lga", "ds"])[list(ANT.values())].sum().reset_index()
obs = []
for a, c in ANT.items():
    g = l.pivot(index="lga", columns="ds", values=c)
    b24 = g.loc[:, g.columns.year == 2024].mean(axis=1); o25 = g.loc[:, g.columns.year == 2025].mean(axis=1)
    obs.append(pd.DataFrame({"lga": b24.index, "antigen": a, "obs_2025_pct_of_2024": (o25 / b24 * 100).values}))
f = pd.read_csv(OUT / "d1_lga_forecast_all.csv").merge(pd.concat(obs), on=["lga", "antigen"], how="left")
f["severity"] = np.where(~f.crosses_80_in_6_12m, "No flag", pd.cut(f.min_forecast_pct_6_12m, [-np.inf, 0, 40, 70, 80],
                         labels=["Unstable forecast (<0%)", "Critical (0-40%)", "Severe (40-70%)", "Warning (70-80%)"]).astype(str))
f["confirmed"] = np.where(f.crosses_80_in_6_12m & (f.obs_2025_pct_of_2024 < 80), "Yes: 2025 doses already below 80% of 2024",
                          np.where(f.crosses_80_in_6_12m, "No: projected dip", "-"))
f["lga_name"] = f.lga.map(clean)
f["obs_2025_pct_of_2024"] = f.obs_2025_pct_of_2024.round(1)
cols = {"zone": "Zone", "state": "State", "lga_name": "LGA", "antigen": "Antigen", "baseline_2024_doses_per_month": "2024 doses per month (baseline)",
        "obs_2025_pct_of_2024": "2025 observed, % of 2024", "min_forecast_pct_6_12m": "Lowest forecast in months 6-12, % of 2024",
        "min_forecast_when": "Month of lowest forecast", "crosses_80_in_6_12m": "Early-warning flag", "severity": "Severity band",
        "confirmed": "Decline already visible in 2025", "first_month_below_80": "First forecast month below 80%", "months_observed": "Months of data"}
F = f[list(cols)].rename(columns=cols).sort_values(["State", "LGA", "Antigen"])
F["Early-warning flag"] = F["Early-warning flag"].map({True: "Yes", False: "No"})
W = F[F["Early-warning flag"] == "Yes"].copy()
W["_o"] = (W["Decline already visible in 2025"].str.startswith("No")).astype(int)
W["_k"] = np.where(W["_o"] == 0, W["2025 observed, % of 2024"], W["Lowest forecast in months 6-12, % of 2024"])
W = W.sort_values(["_o", "_k"]).drop(columns=["_o", "_k"])
S = f.groupby(["zone", "state"]).agg(LGAs=("lga", "nunique"), Forecasts=("antigen", "size"), Alerts=("crosses_80_in_6_12m", "sum"),
                                     Confirmed=("confirmed", lambda x: int(x.str.startswith("Yes").sum()))).reset_index()
fl = f[f.crosses_80_in_6_12m]
S = S.merge(fl.groupby("state").lga.nunique().rename("LGAs with any alert"), on="state", how="left").fillna(0)
for a in ANT:
    S = S.merge(fl[fl.antigen == a].groupby("state").size().rename(f"{a} alerts"), on="state", how="left").fillna(0)
S = S.rename(columns={"zone": "Zone", "state": "State", "Confirmed": "Alerts with decline visible in 2025"}).sort_values(["Zone", "State"])
N = pd.read_csv(OUT / "d1_national_at_risk_summary.csv").rename(columns={"antigen": "Antigen", "base_2024_doses_per_month": "2024 doses per month",
    "min_forecast_pct": "Lowest forecast, % of 2024", "min_when": "Month of lowest", "crosses_80_in_6_12m": "Below 80% in months 6-12",
    "min_pct_in_window": "Lowest in months 6-12, % of 2024"})
A = pd.read_csv(OUT / "D1_additional_antigen_forecast_summary.csv")
CV = pd.read_csv(ROOT / "figures" / "data" / "D1_F2_cohort_coverage_vs_ndhs.csv").rename(columns={"antigen": "Antigen",
    "admin_coverage_2026_pct": "Administrative coverage 2026, % (forecast doses / eligible cohort)", "ndhs_2024_pct": "NDHS 2024 survey coverage, %",
    "matched_states": "States"})
c = pd.read_csv(DATA / "dhis2_data_consolidated.csv", thousands=",", low_memory=False)
gz = c[c.lga.str.contains("Guzamala", case=False)].copy()
gz["ds"] = pd.to_datetime(gz.period.astype(str).str.strip(), format="%b-%y"); gz = gz.sort_values("ds")
cnt = [k for k in gz.columns if k.endswith("_count")]
G = gz[["state", "lga", "period"] + cnt].rename(columns={"state": "State", "lga": "DHIS2 organisation unit", "period": "Month"})
G.insert(3, "Any antigen reported", np.where(gz[cnt].notna().any(axis=1), "Yes", "No"))
G.insert(4, "Penta1 reported", np.where(gz.penta_1_count.notna(), "Yes", "No"))
summ = json.loads((OUT / "d1_summary.json").read_text())
about = pd.DataFrame([
    ("What this workbook contains", "Domain 1 early-warning forecasts of routine immunization doses for the four tracer antigens (BCG, Penta1, Penta3, "
     "Measles1) in every local government area (LGA) with a 2024 baseline, plus national forecasts for nine additional antigens."),
    ("Method", "Prophet time-series model per LGA and antigen (yearly and semi-annual seasonality, 95% prediction interval) fitted to DHIS2 monthly "
     "doses for January 2021 to December 2025 and projected 18 months. Each forecast is expressed as a percent of the same LGA's 2024 mean monthly "
     "doses. An early-warning flag is raised when the forecast falls below 80% in months 6 to 12 (June to December 2026)."),
    ("What the flag means", "Doses delivered are projected to fall below four-fifths of what the LGA delivered in 2024. It is a decline tripwire, "
     "not a coverage figure: no child population enters the calculation."),
    ("Coverage", f"{summ['forecasts_fitted']:,} LGA-antigen forecasts for {summ['lgas_with_any_forecast']} LGAs; {summ['alerts']:,} early-warning flags "
     f"in {summ['lgas_with_any_alert']} LGAs. Guzamala (Borno) has no forecast because it reported no doses in 2024 (see 'Guzamala check')."),
    ("Priority", "The worklist lists flags where the decline is already visible in 2025 data first (2025 observed doses below 80% of 2024), "
     "then projected dips."),
    ("Unstable forecasts", "Forecasts below 0% indicate an erratic LGA series (for example, a sudden change in reporting) rather than a real "
     "projection; verify the LGA's DHIS2 data before acting."),
    ("Data", "NPHCDA DHIS2 routine immunization data by LGA and month, 2021-2025. Katsina, Kogi, Kwara, Lagos and Niger have no rows before "
     "January 2023 in this export. NDHS 2024 (survey coverage); NPC under-five projection 2024 (eligible cohort)."),
    ("Status", "Model estimates for programme planning."),
], columns=["Item", "Description"])
sheets = [("About", about), ("National tracers", N), ("Additional antigens", A), ("Coverage vs survey", CV), ("Worklist (flagged)", W),
          ("LGA forecasts (all)", F), ("State summary", S), ("Guzamala check", G)]
with pd.ExcelWriter(WB, engine="openpyxl") as xw:
    for name, df in sheets:
        df.to_excel(xw, sheet_name=name, index=False, startrow=2)
wb = load_workbook(WB)
titles = {"About": "Domain 1: antigen early-warning forecasts, 2026", "National tracers": "National forecasts, four tracer antigens (% of 2024 level)",
          "Additional antigens": "National forecasts, nine additional antigens", "Coverage vs survey": "Administrative coverage of the eligible cohort vs NDHS 2024",
          "Worklist (flagged)": "Early-warning worklist: LGA-antigen flags, decline visible in 2025 first", "LGA forecasts (all)": "All LGA-antigen forecasts",
          "State summary": "Early-warning flags by state", "Guzamala check": "Guzamala LGA (Borno): DHIS2 records by month, 2021-2025 (for verification)"}
for ws in wb.worksheets:
    ws["A1"] = titles[ws.title]; ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    hdr = ws[3]
    for cell in hdr:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 42
    for j, col in enumerate(ws.iter_cols(min_row=3, max_row=min(ws.max_row, 400)), 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, w * 0.9), 60 if ws.title != "About" else 130)
    ws.freeze_panes = "A4"
    if ws.max_row > 4 and ws.title not in ("About",):
        ref = f"A3:{get_column_letter(ws.max_column)}{ws.max_row}"
        t = Table(displayName=ws.title.replace(" ", "").replace("(", "").replace(")", "").replace("-", "")[:30], ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)
    if ws.title == "About":
        for r in ws.iter_rows(min_row=4):
            r[1].alignment = Alignment(wrap_text=True, vertical="top"); r[0].font = Font(bold=True, color=NAVY)
    if ws.title == "Guzamala check":
        red = PatternFill("solid", fgColor="FBE5D6")
        for r in ws.iter_rows(min_row=4):
            if r[4].value == "No":
                r[4].fill = red
wb.save(WB)
G.to_csv(ROOT / "outputs" / "Guzamala_Borno_DHIS2_records_2021_2025.csv", index=False)
print("saved", WB.name, {n: len(df) for n, df in sheets})
print("worklist confirmed:", int((W["Decline already visible in 2025"].str.startswith("Yes")).sum()))
